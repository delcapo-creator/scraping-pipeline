"""Write datasets to CSV, JSON, SQLite and Excel (.xlsx)."""
from __future__ import annotations

import csv
import json
import logging
import sqlite3
from pathlib import Path

log = logging.getLogger(__name__)


def _flat(v):
    return "|".join(map(str, v)) if isinstance(v, list) else v


def write_csv(rows: list[dict], path: Path) -> None:
    if not rows:
        return
    with open(path, "w", newline="", encoding="utf-8-sig") as f:   # utf-8-sig: opens cleanly in Excel
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows({k: _flat(v) for k, v in r.items()} for r in rows)


def write_json(rows: list[dict], path: Path) -> None:
    path.write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8")


def write_sqlite(rows: list[dict], path: Path, table: str) -> None:
    if not rows:
        return
    cols = list(rows[0])
    con = sqlite3.connect(path)
    try:
        con.execute(f'DROP TABLE IF EXISTS "{table}"')
        con.execute(f'CREATE TABLE "{table}" ({", ".join(f"{c!r}" for c in cols)})')
        con.executemany(f'INSERT INTO "{table}" VALUES ({", ".join("?" * len(cols))})',
                        [[_flat(r.get(c)) for c in cols] for r in rows])
        con.commit()
    finally:
        con.close()


def write_xlsx(rows: list[dict], path: Path, sheet: str = "data") -> None:
    """Real Excel file: opens with proper columns on any locale (a CSV opens in one column
    on French/German Excel, which expects ';'). Numbers stay numbers, header is bold and
    frozen, filters are on, column widths fit the content."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    if not rows:
        return
    cols = list(rows[0])
    wb = Workbook()
    ws = wb.active
    ws.title = sheet[:31]
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([_flat(r.get(c)) for c in cols])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, c in enumerate(cols, 1):
        longest = max([len(str(c))] + [len(str(_flat(r.get(c)) or "")) for r in rows[:500]])
        ws.column_dimensions[get_column_letter(i)].width = min(60, longest + 2)
    wb.save(path)


def _xlsx_available() -> bool:
    try:
        import openpyxl  # noqa: F401
        return True
    except ImportError:
        return False


def write_all(rows: list[dict], out_dir: Path, name: str, formats: list[str]) -> list[Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    written = []
    for fmt in formats:
        if fmt == "xlsx" and not _xlsx_available():
            log.warning("xlsx skipped: run  pip install openpyxl")
            continue
        p = out_dir / f"{name}.{ 'db' if fmt == 'sqlite' else fmt }"
        {"csv": lambda: write_csv(rows, p), "json": lambda: write_json(rows, p),
         "sqlite": lambda: write_sqlite(rows, p, name), "xlsx": lambda: write_xlsx(rows, p, name)}[fmt]()
        written.append(p)
    return written
